import importlib.util
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / ".agents" / "skills" / "document-export" / "scripts" / "export_documents.py"
SPEC = importlib.util.spec_from_file_location("document_export", SCRIPT)
exporter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(exporter)


class DocumentExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.document = {"title": "Candidate & <Profile>", "subtitle": "candidate@example.test",
                         "sections": [{"heading": "Experience", "paragraphs": ["Built <safe> reports & dashboards."],
                                       "bullets": ["Measured outcomes"],
                                       "table": {"headers": ["Skill", "Evidence"], "rows": [["Python", "Report"]]}}]}

    def test_probe_is_stdlib_safe_and_non_mutating(self):
        result = exporter.probe()
        self.assertEqual(result["status"], "OK")
        self.assertFalse(result["writes"])
        self.assertIn("docx", result["dependencies"])

    def test_generic_default_keeps_personal_reference_as_explicit_option(self):
        self.assertEqual(exporter.export_document.__kwdefaults__["template"], "modern-cn")
        self.assertEqual(exporter.export_document.__kwdefaults__["page_policy"], "a4-1-3")
        self.assertEqual(exporter.TEMPLATES["reference"]["label"], "个人参考模板")
        self.assertEqual(exporter.TEMPLATES["reference"]["margin"], 1.35)

    def test_historical_draft_keeps_source_correction_diagnostics(self):
        from oriward import experience_library as lib, services, sources
        with patch.dict(os.environ, {"ORIWARD_DB": str(self.root / "isolated.sqlite3"),
                                     "ORIWARD_OUTPUT_DIR": str(self.root / "resumes")}):
            services.save_profile({"display_name": "Fictional Candidate"})
            eid = lib.create({"event": "Fictional coursework", "actions": ["Built a report"], "outcomes": [],
                              "evidence": [{"claim": "Built a report", "source": "self-report", "status": "user_stated"}],
                              "approved_for_resume": True})
            jid = sources.import_records([{"source": "fictional", "source_id": "1", "title": "Intern",
                "company": "Fictional", "description": "Reports", "url": "https://example.invalid/1"}])["job_ids"][0]
            resume = services.build_resume(jid, [eid])
            lib.exclude(eid, reason="Candidate excluded it")
            document, provenance = exporter._load_resume(resume["id"])
            self.assertTrue(provenance["source_validity"]["needs_review"])
            self.assertEqual(len(provenance["markdown_sha256"]), 64)
            self.assertIn("Fictional coursework", exporter._all_text(document))

    def test_schema_rejects_unknown_and_malformed_values(self):
        with self.assertRaisesRegex(exporter.ExportError, "unknown"):
            exporter.export_document(document={"sections": [], "extra": True}, output_dir=self.root / "x", formats=("docx",))
        with self.assertRaisesRegex(exporter.ExportError, "array of strings"):
            exporter.export_document(document={"sections": [{"paragraphs": [7]}]}, output_dir=self.root / "y", formats=("docx",))

    def test_existing_output_is_never_touched(self):
        target = self.root / "already-there"
        target.mkdir()
        sentinel = target / "keep.txt"
        sentinel.write_text("do not alter", encoding="utf-8")
        with self.assertRaisesRegex(exporter.ExportError, "OUTPUT_COLLISION"):
            exporter.export_document(document=self.document, output_dir=target, formats=("docx",))
        self.assertEqual(sentinel.read_text(encoding="utf-8"), "do not alter")

    def test_word_renderer_uses_absolute_paths_and_exports_pdf(self):
        from unittest.mock import patch
        source = self.root / "input.docx"
        target = self.root / "nested" / "output.pdf"
        target.parent.mkdir()
        source.write_bytes(b"docx")
        observed = {}

        def run(args, **kwargs):
            observed["args"] = args
            observed["script"] = Path(args[4]).read_text(encoding="utf-8")
            Path(args[-1]).write_bytes(b"pdf")

        with patch.object(exporter.subprocess, "run", side_effect=run):
            exporter._render_docx_pdf(source, target, "word", ("word", self.root / "WINWORD.EXE"))
        self.assertTrue(Path(observed["args"][-2]).is_absolute())
        self.assertTrue(Path(observed["args"][-1]).is_absolute())
        self.assertIn("ExportAsFixedFormat($OutputPath,17)", observed["script"])
        self.assertNotIn("SaveAs($OutputPath,17)", observed["script"])
        self.assertTrue(target.is_file())

    def test_missing_glyph_font_is_refused_before_output(self):
        vera = exporter._vera_path()
        if not vera:
            self.skipTest("ReportLab Vera is unavailable")
        with self.assertRaisesRegex(exporter.ExportError, "FONT_COVERAGE"):
            exporter.export_document(document={"title": "中文", "sections": [{"paragraphs": ["中文"]}]},
                                     output_dir=self.root / "cjk", formats=("docx",), font=vera)
        self.assertFalse((self.root / "cjk").exists())

    def test_file_created_after_output_mkdir_is_not_overwritten(self):
        target = self.root / "race-output"
        original_mkdir = Path.mkdir
        def create_racing_file(path, *args, **kwargs):
            result = original_mkdir(path, *args, **kwargs)
            if path == target:
                (target / "document.docx").write_bytes(b"third-party")
            return result
        # Exercise the exclusive-write race independently of system fonts and
        # optional renderers: generation must never start for this collision.
        with patch.object(Path, "mkdir", new=create_racing_file), \
                patch.object(exporter, "_choose_font", return_value=self.root / "fixture.ttf"), \
                patch.object(exporter, "_font_family", return_value="Fixture"), \
                patch.object(exporter, "_write_docx") as write_docx:
            with self.assertRaisesRegex(exporter.ExportError, "OUTPUT_COLLISION.*document.docx"):
                exporter.export_document(document=self.document, output_dir=target, formats=("docx",))
        write_docx.assert_not_called()
        self.assertEqual((target / "document.docx").read_bytes(), b"third-party")

    def test_real_generated_resume_keeps_summary_headline_experience_order_and_provenance(self):
        try:
            import docx
            import reportlab  # noqa: F401 -- optional dependency availability probe
            from pypdf import PdfReader
            from oriward import services as s, sources
        except ImportError:
            self.skipTest("optional document libraries or Oriward unavailable")
        with patch.dict(os.environ, {"ORIWARD_DB": str(self.root / "db.sqlite3"),
                                     "ORIWARD_OUTPUT_DIR": str(self.root / "resumes")}, clear=False):
            s.save_profile({"display_name": "Fictional Candidate", "contact": {"email": "fictional@example.test"},
                            "summary_verified": "Verified profile summary"})
            experience_ids = []
            for event, bullet in (("First project", "Built first report"), ("Second project", "Built second report")):
                experience_ids.append(s.add_object("experiences", {"event": event, "actions": [bullet], "outcomes": [],
                    "evidence": [{"claim": bullet, "source": "fixture", "status": "user_confirmed"}],
                    "approved_for_resume": True}, ["event", "actions", "outcomes", "evidence", "approved_for_resume"]))
            job = sources.import_records([{ "source": "fixture", "source_id": "real-export", "title": "Analyst", "company": "Fixture",
                "description": "Test", "availability": "listed", "url": "https://example.invalid/real-export"}])["job_ids"][0]
            resume = s.build_resume(job, experience_ids, headline="Target role")
            resume_path = Path(resume["path"])
            original = resume_path.read_text(encoding="utf-8")
            resume_path.write_text(original.replace("### First project", "### First project\n\n- Edited locally after factual approval"),
                                    encoding="utf-8")
            # Generated resumes have Chinese section labels even with English
            # fixture facts. A host without an embeddable CJK font must report
            # that limitation rather than be marked as a broken DOCX backend.
            try:
                exporter._choose_font(["个人简介 相关经历 申请方向：Target role"], None)
            except exporter.ExportError as exc:
                if str(exc).startswith("FONT_COVERAGE"):
                    self.skipTest("no local embeddable CJK font for generated resume labels")
                raise
            output = self.root / "resume-output"
            result = exporter.export_document(resume_id=resume["id"], output_dir=output)
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["source"]["resume_id"], resume["id"])
        self.assertEqual(manifest["source"]["job_id"], job)
        self.assertEqual(manifest["source"]["omitted_internal_footer"], "*来源记录：" in original)
        self.assertTrue(manifest["exported_at"])
        docx_text = "\n".join(p.text for p in docx.Document(output / "document.docx").paragraphs)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(output / "document.pdf").pages)
        for expected in ("Verified profile summary", "申请方向：Target role", "First project", "Edited locally after factual approval",
                         "Built first report", "Second project", "Built second report"):
            self.assertIn(expected, docx_text)
        for text in (docx_text, pdf_text):
            for expected in ("Verified profile summary", "申请方向：Target role", "First project", "Edited locally after factual approval",
                             "Built first report", "Second project", "Built second report"):
                # PDF text extraction may insert spacing between CJK and Latin runs.
                self.assertIn("".join(expected.split()), "".join(text.split()))
            self.assertLess(text.index("First project"), text.index("Built first report"))
            self.assertLess(text.index("Built first report"), text.index("Second project"))
            self.assertLess(text.index("Second project"), text.index("Built second report"))
            self.assertNotIn("来源记录", text)
        self.assertEqual(result["status"], "GENERATED_PENDING_VISUAL_REVIEW")

    def test_pending_resume_revision_exports_for_local_reading(self):
        try:
            from oriward import feedback, services as s, sources
        except ImportError:
            self.skipTest("Oriward package unavailable")
        with patch.dict(os.environ, {"ORIWARD_DB": str(self.root / "db.sqlite3"),
                                     "ORIWARD_OUTPUT_DIR": str(self.root / "resumes")}, clear=False):
            s.save_profile({"display_name": "Fictional Candidate"})
            exp = s.add_object("experiences", {"event": "Demo", "actions": ["Built report"], "outcomes": [],
                "evidence": [{"claim": "Built report", "source": "fixture", "status": "user_confirmed"}],
                "approved_for_resume": True}, ["event", "actions", "outcomes", "evidence", "approved_for_resume"])
            job = sources.import_records([{ "source": "fixture", "source_id": "pending", "title": "Analyst", "company": "Fixture",
                "description": "Test", "availability": "listed", "url": "https://example.invalid/pending"}])["job_ids"][0]
            base = s.build_resume(job, [exp])
            revision = feedback.revise_resume({"base_resume_id": base["id"], "reason": "Rewrite fixture wording", "feedback_ids": [],
                "sections": [{"experience_id": exp, "bullets": [{"text": "Created report", "source_texts": ["Built report"]}]}]})
            # Simulate a legacy pending version without changing its stored
            # status: local preview must not turn that status into a gate.
            from oriward.db import connect
            with connect() as db:
                db.execute("UPDATE resume_versions SET factual_status='pending' WHERE resume_id=?", (revision["id"],))
            document, provenance = exporter._load_resume(revision["id"])
            self.assertEqual(provenance["factual_status"], "pending")
            self.assertIn("Created report", exporter._all_text(document))
            try:
                import docx
                exporter._choose_font(exporter._all_text(document), None)
            except ImportError:
                self.skipTest("python-docx unavailable")
            except exporter.ExportError as exc:
                self.skipTest(f"No local font: {exc}")
            output = self.root / "pending-output"
            manifest = exporter.export_document(resume_id=revision["id"], output_dir=output, formats=("docx",))
            self.assertEqual(manifest["source"]["factual_status"], "pending")
            text = "\n".join(p.text for p in docx.Document(output / "document.docx").paragraphs)
            self.assertIn("Created report", text)
            for internal in ("来源记录", "经历 ID", "factual_status", "pending"):
                self.assertNotIn(internal, text)

    def test_structured_input_cannot_pose_as_a_resume(self):
        resume_like = {"title": "Fictional Candidate", "sections": [
            {"heading": "相关经历", "bullets": ["5 年数据分析经验"]}]}
        with self.assertRaisesRegex(exporter.ExportError, "RESUME_REQUIRES_ID"):
            exporter.export_document(document=resume_like, output_dir=self.root / "fake-resume",
                                     formats=("docx",))
        self.assertFalse((self.root / "fake-resume").exists())

    def test_resume_parser_keeps_two_experiences_in_content_order(self):
        source = "\n".join(["# Candidate", "", "## 相关经历", "", "### First project", "", "- First evidence", "",
                              "### Second project", "", "- Second evidence", "", "---", "",
                              "*来源记录：岗位 1（Fixture / Analyst）；经历 ID 1, 2。*",
                              "*发布/投递前请人工核对内容；未经证实的推断不应当写入简历。*", ""])
        document, _ = exporter._parse_resume_markdown(source)
        ordered = []
        for section in document["sections"]:
            ordered.append(section["heading"])
            ordered.extend(section["paragraphs"])
            ordered.extend(section["bullets"])
        self.assertEqual(ordered, ["相关经历", "First project", "First evidence", "Second project", "Second evidence"])

    def test_resume_parser_accepts_clean_adult_draft_without_internal_footer(self):
        source = "\n".join(["# Fictional Candidate", "", "## 工作经历", "",
                            "### Service company / Operations", "*2016–2022*", "- Coordinated schedules", "",
                            "## 经营经历", "", "### Community shop", "*2022–2025*", "- Managed inventory", "",
                            "## 其他经历", "", "### Community activity", "- Organized volunteers", ""])
        document, footer = exporter._parse_resume_markdown(source)
        self.assertEqual(footer, [])
        self.assertEqual([s["heading"] for s in document["sections"]],
                         ["工作经历", "Service company / Operations", "经营经历", "Community shop",
                          "其他经历", "Community activity"])
        self.assertEqual(document["sections"][1]["meta"], "2016–2022")
        self.assertEqual(document["sections"][3]["meta"], "2022–2025")

    def test_resume_parser_preserves_section_hierarchy_dates_and_profile_order(self):
        source = "\n".join(["# Candidate", "", "candidate@example.test", "*Product intern*", "",
                              "## 教育经历", "", "河海大学金融硕士", "", "## 相关经历", "",
                              "### Oriward", "", "*2026年，持续迭代*", "", "- 项目主页：[Oriward GitHub](https://example.invalid/project)", "",
                              "---", "", "*来源记录：岗位 1（Fixture / Analyst）；经历 ID 1。*",
                              "*发布/投递前请人工核对内容；未经证实的推断不应当写入简历。*", ""])
        document, _ = exporter._parse_resume_markdown(source)
        self.assertEqual(document["subtitle"], "candidate@example.test\nProduct intern")
        self.assertEqual(document["sections"][0]["level"], 1)
        self.assertEqual(document["sections"][2]["level"], 2)
        self.assertEqual(document["sections"][2]["meta"], "2026年，持续迭代")
        self.assertIn("[Oriward GitHub](https://example.invalid/project)", document["sections"][2]["bullets"][0])

    def test_paired_outputs_read_same_key_text_and_keep_html_literal(self):
        try:
            import docx
            import reportlab  # noqa: F401 -- optional dependency availability probe
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("DOCX/PDF optional libraries unavailable")
        document = self.document
        document["sections"][0]["bullets"] = ["Project: [Oriward GitHub](https://example.invalid/project)"]
        output = self.root / "generated"
        result = exporter.export_document(document=document, output_dir=output)
        self.assertEqual(result["status"], "GENERATED_PENDING_VISUAL_REVIEW")
        docx_text = "\n".join(p.text for p in docx.Document(output / "document.docx").paragraphs)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(output / "document.pdf").pages)
        for text in (docx_text, pdf_text):
            self.assertIn("Candidate & <Profile>", text)
            self.assertIn("Built <safe> reports & dashboards.", text)
        relationships = docx.Document(output / "document.docx").part.rels
        self.assertTrue(any(rel.reltype.endswith("/hyperlink") and rel.target_ref == "https://example.invalid/project"
                            for rel in relationships.values()))
        self.assertEqual(json.loads((output / "manifest.json").read_text(encoding="utf-8"))["status"],
                         "GENERATED_PENDING_VISUAL_REVIEW")

    def test_long_table_starts_on_page_with_available_space_after_its_heading(self):
        try:
            from pypdf import PdfReader
            import reportlab  # noqa: F401 -- optional dependency availability probe
        except ImportError:
            self.skipTest("ReportLab/PDF reader unavailable")
        document = {"title": "Table pagination", "sections": [
            {"heading": "Intro", "paragraphs": ["Earlier content keeps enough room for a table start."] * 18},
            {"heading": "Long table", "table": {"headers": ["Number", "Evidence"],
             "rows": [[f"Row {index}", "A deliberately longer value which wraps across the available table width."]
                      for index in range(1, 90)]}}]}
        output = self.root / "long-table"
        exporter.export_document(document=document, output_dir=output, formats=("pdf",))
        reader = PdfReader(output / "document.pdf")
        self.assertGreaterEqual(len(reader.pages), 2)
        first_page = reader.pages[0].extract_text() or ""
        self.assertIn("Long table", first_page)
        self.assertIn("Row 1", first_page)


if __name__ == "__main__":
    unittest.main()
